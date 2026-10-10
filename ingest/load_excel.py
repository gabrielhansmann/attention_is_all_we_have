"""Load the provided Excel files into Supabase.

  *_feedback_hackathon.xlsx  -> feedback_raw   (one row per comment, comment embedded with OpenAI)
  sales_volumes.xlsx         -> sales_volume   (one row per line/market/year)
  *customer_studies.xlsx     -> study_response (cross-tab cells unpacked into rows)

Every row keeps its sheet and row (or cell) so it can be cited. Re-running replaces a file's rows.
Usage:  python ingest/load_excel.py [feedback|sales|studies ...]   (default: all)
Needs OPENAI_API_KEY and SUPABASE_DB_URL in .env.
"""
import os
import sys
from pathlib import Path

import openpyxl
import psycopg
from dotenv import load_dotenv
from openai import OpenAI

load_dotenv(".env")
DATA = Path("data")
EMBED_MODEL = "text-embedding-3-small"
BATCH = 200
MAX_EMBED_CHARS = 6000


def vec(v):
    return "[" + ",".join(f"{x:.7f}" for x in v) + "]"


def txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def source_id(conn, name, vehicle_line=None):
    row = conn.execute("select id from public.source where uri = %s", (name,)).fetchone()
    if row:
        return row[0]
    return conn.execute(
        "insert into public.source (type, uri, vehicle_line, title) values ('internal_xlsx', %s, %s, %s) returning id",
        (name, vehicle_line, name),
    ).fetchone()[0]


def load_feedback(conn, client):
    for path in sorted(DATA.glob("*_feedback_hackathon.xlsx")):
        line = path.name.split("_")[0]
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [(i, r) for i, r in enumerate(ws.iter_rows(values_only=True), 1) if i > 1 and r[0]]
        sid = source_id(conn, path.name, line)
        with conn.transaction():
            conn.execute("delete from public.feedback_raw where source_id = %s", (sid,))
            for b in range(0, len(rows), BATCH):
                batch = rows[b:b + BATCH]
                emb = client.embeddings.create(
                    model=EMBED_MODEL,
                    input=[(txt(r[8]) or "")[:MAX_EMBED_CHARS] or "(empty)" for _, r in batch],
                ).data
                params = [
                    (sid, ws.title, i, txt(r[0]), txt(r[1]), txt(r[2]), txt(r[3]), txt(r[4]), txt(r[5]),
                     txt(r[6]), txt(r[7]), txt(r[8]), txt(r[9]), txt(r[10]), txt(r[11]),
                     vec(e.embedding) if txt(r[8]) else None)
                    for (i, r), e in zip(batch, emb)
                ]
                with conn.cursor() as cur:
                    cur.executemany(
                        "insert into public.feedback_raw (source_id, sheet, row_no, ext_id, feedback_source, brand, "
                        "sales_description, country, derivate, engine_type, feedback_type, comment, vfc_label, "
                        "vfc_level2, vfc_level3, embedding) values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s::extensions.vector)",
                        params,
                    )
                print(f"  {path.name}: {min(b + BATCH, len(rows))}/{len(rows)}", flush=True)
        print(f"{path.name}: {len(rows)} comments loaded")


def load_sales(conn):
    path = DATA / "sales_volumes.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)
    sid = source_id(conn, path.name)
    n = 0
    with conn.transaction():
        conn.execute("delete from public.sales_volume where source_id = %s", (sid,))
        for ws in wb.worksheets:
            for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                if not r[0]:
                    continue
                for year, val in zip((2024, 2025, 2030), r[2:5]):
                    conn.execute(
                        "insert into public.sales_volume (source_id, sheet, row_no, market, market_code, year, volume, is_forecast) "
                        "values (%s,%s,%s,%s,%s,%s,%s,%s)",
                        (sid, ws.title, i, r[0], r[1], year, int(val), year > 2025),
                    )
                    n += 1
    print(f"{path.name}: {n} rows loaded")


def parse_cn_eu(ws):
    """Rows: question label row, then a 'Mean' row with one value per model column."""
    countries, col_country, cur = {}, {}, None
    for c in range(2, ws.max_column + 1):
        cur = txt(ws.cell(2, c).value) or cur
        col_country[c] = cur
    models = {c: txt(ws.cell(3, c).value) for c in range(2, ws.max_column + 1) if txt(ws.cell(3, c).value)}
    question = txt(ws.cell(3, 1).value)
    for r in range(4, ws.max_row + 1):
        label = txt(ws.cell(r, 1).value)
        if label == "Mean":
            for c, model in models.items():
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)):
                    yield ws.cell(r, c).coordinate, col_country[c], model, question, "mean", None, v
        elif label:
            question = label


def parse_us(ws):
    """Blocks: question label row, then 'Sample total' and seven answer rows with shares per model."""
    models = {c: txt(ws.cell(1, c).value) for c in range(3, ws.max_column + 1) if txt(ws.cell(1, c).value)}
    question = None
    for r in range(1, ws.max_row + 1):
        label, kind = txt(ws.cell(r, 1).value), txt(ws.cell(r, 2).value)
        if label and not kind:
            question = label
        elif label and kind and label != "Sample total":
            for c, model in models.items():
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)):
                    yield ws.cell(r, c).coordinate, "US", model, question, "share", label, v


def load_studies(conn):
    path = next(DATA.glob("*customer_studies.xlsx"))
    wb = openpyxl.load_workbook(path, data_only=True)
    sid = source_id(conn, path.name)
    n = 0
    with conn.transaction():
        conn.execute("delete from public.study_response where source_id = %s", (sid,))
        for ws in wb.worksheets:
            parser = parse_cn_eu if ws.title.startswith("CN_EU") else parse_us
            for cell, market, model, question, metric, option, value in parser(ws):
                conn.execute(
                    "insert into public.study_response (source_id, sheet, cell_ref, market, model, question, metric, answer_option, value) "
                    "values (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                    (sid, ws.title, cell, market, model, question, metric, option, value),
                )
                n += 1
    print(f"{path.name}: {n} cells loaded")


if __name__ == "__main__":
    what = set(sys.argv[1:]) or {"feedback", "sales", "studies"}
    # autocommit so each `with conn.transaction()` block is a real, committed transaction
    conn = psycopg.connect(os.environ["SUPABASE_DB_URL"], autocommit=True)
    if "sales" in what:
        load_sales(conn)
    if "studies" in what:
        load_studies(conn)
    if "feedback" in what:
        load_feedback(conn, OpenAI())
    conn.close()
