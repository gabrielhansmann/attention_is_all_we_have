"""Load the provided Excel files into Supabase.

  *_feedback_hackathon.xlsx  -> feedback_raw   (one row per comment, comment embedded with OpenAI)
  sales_volumes.xlsx         -> sales_volume   (one row per line/market/year)
  *customer_studies.xlsx     -> study_response (cross-tab cells unpacked into rows)

Every row keeps its sheet and row (or cell) so it can be cited. Re-running replaces a file's rows.
The ingest_* functions load one file each; the upload endpoint (app.ingest.upload) uses them too.
Usage:  python -m app.ingest.load_excel [feedback|sales|studies ...]   (default: all)
Needs OPENAI_API_KEY and SUPABASE_DB_URL in .env.
"""
import os
import re
import sys

import openpyxl
import psycopg
from app.config import DATA_DIR, load_env
from openai import OpenAI

load_env()
DATA = DATA_DIR
EMBED_MODEL = "text-embedding-3-small"
BATCH = 200
MAX_EMBED_CHARS = 6000
FEEDBACK_COLS = 12
SALES_YEARS = (2024, 2025, 2030)


def vec(v):
    return "[" + ",".join(f"{x:.7f}" for x in v) + "]"


def txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


class LayoutError(ValueError):
    """The workbook does not have the layout this loader reads, or holds no usable rows."""


def source_id(conn, name, vehicle_line=None, type="internal_xlsx", title=None):
    row = conn.execute("select id from public.source where uri = %s", (name,)).fetchone()
    if row:
        if vehicle_line:
            conn.execute("update public.source set vehicle_line = %s where id = %s", (vehicle_line, row[0]))
        return row[0]
    return conn.execute(
        "insert into public.source (type, uri, vehicle_line, title) values (%s, %s, %s, %s) returning id",
        (type, name, vehicle_line, title or name),
    ).fetchone()[0]


# The same vocabulary the knowledge-layer migration backfilled (20261011_knowledge_layer.sql), applied to the
# rows of one file so newly loaded evidence is found by the same filters.
FEEDBACK_NORMALIZE = """
update public.feedback_raw f set
  vehicle_line = coalesce(f.derivate, %(line)s),
  market_code = (select c.market_code from public.country_market c where c.country_code = f.country),
  domain_id = (select o.id from public.ontology o
               where o.kind = 'domain' and o.parent_id is not null and o.label = f.vfc_level2 limit 1),
  polarity = case f.feedback_type when 'Likes' then 'positive' when 'Difficult to Use' then 'negative'
                                  when 'Defect' then 'negative' when 'Wants' then 'request' end,
  powertrain = case f.engine_type when 'ICE' then 'ICE' when 'PHEV' then 'PHEV'
                                  when 'BEVE' then 'BEV' when 'BEV' then 'BEV' end
where f.source_id = %(sid)s"""
SALES_NORMALIZE = """
update public.sales_volume set vehicle_line = case
  when sheet = 'G60_G68' and market_code = 'CN' then 'G68' when sheet = 'G60_G68' then 'G60' else sheet end
where source_id = %s"""
STUDY_NORMALIZE = """
update public.study_response r set
  vehicle_line = (select o.key from public.ontology o
                  where o.kind = 'vehicle_line' and r.model like '%%' || o.key || '%%' order by o.key limit 1),
  market_code = case r.market when 'China' then 'CN' when 'EU' then 'EU' when 'US' then 'US' end
where r.source_id = %s"""


def ingest_feedback(conn, client, path, line=None):
    """One feedback export -> feedback_raw. Returns (comments loaded, rows skipped for having no ID)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    sheet, rows, skipped = ws.title, [], 0
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        r = tuple(r) + (None,) * (FEEDBACK_COLS - len(r))
        if txt(r[0]):
            rows.append((i, r))
        elif any(txt(v) for v in r):
            skipped += 1
    wb.close()
    if not rows:
        raise LayoutError("no feedback rows found")

    params = []
    for b in range(0, len(rows), BATCH):
        batch = rows[b:b + BATCH]
        emb = client.embeddings.create(
            model=EMBED_MODEL,
            input=[(txt(r[8]) or "")[:MAX_EMBED_CHARS] or "(empty)" for _, r in batch],
        ).data
        params += [
            (sheet, i, txt(r[0]), txt(r[1]), txt(r[2]), txt(r[3]), txt(r[4]), txt(r[5]),
             txt(r[6]), txt(r[7]), txt(r[8]), txt(r[9]), txt(r[10]), txt(r[11]),
             vec(e.embedding) if txt(r[8]) else None)
            for (i, r), e in zip(batch, emb)
        ]
        print(f"  {path.name}: {min(b + BATCH, len(rows))}/{len(rows)} embedded", flush=True)

    with conn.transaction():
        sid = source_id(conn, path.name, line)
        conn.execute("delete from public.feedback_raw where source_id = %s", (sid,))
        with conn.cursor() as cur:
            cur.executemany(
                "insert into public.feedback_raw (source_id, sheet, row_no, ext_id, feedback_source, brand, "
                "sales_description, country, derivate, engine_type, feedback_type, comment, vfc_label, "
                "vfc_level2, vfc_level3, embedding) values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s::extensions.vector)",
                [(sid, *p) for p in params],
            )
        conn.execute(FEEDBACK_NORMALIZE, {"sid": sid, "line": line})
    print(f"{path.name}: {len(rows)} comments loaded")
    return len(rows), skipped


def load_feedback(conn, client):
    for path in sorted(DATA.glob("*_feedback_hackathon.xlsx")):
        ingest_feedback(conn, client, path, path.name.split("_")[0])


def _years(ws):
    """The year of each volume column, read from the header; the provided layout is the fallback."""
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    found = [re.search(r"(?<!\d)(20\d{2})(?!\d)", str(v or "")) for v in header[2:5]]
    return tuple(int(m.group(1)) for m in found) if len(found) == 3 and all(found) else SALES_YEARS


def ingest_sales(conn, path):
    """One sales workbook (a sheet per vehicle line) -> sales_volume. Returns (rows loaded, cells skipped)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    n = skipped = 0
    with conn.transaction():
        sid = source_id(conn, path.name)
        conn.execute("delete from public.sales_volume where source_id = %s", (sid,))
        for ws in wb.worksheets:
            years = _years(ws)
            for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                if not r or not txt(r[0]):
                    continue
                for year, val in zip(years, r[2:5]):
                    if not txt(r[1]) or isinstance(val, bool) or not isinstance(val, (int, float)):
                        skipped += 1
                        continue
                    conn.execute(
                        "insert into public.sales_volume (source_id, sheet, row_no, market, market_code, year, volume, is_forecast) "
                        "values (%s,%s,%s,%s,%s,%s,%s,%s)",
                        (sid, ws.title, i, txt(r[0]), txt(r[1]), year, int(val), year > 2025),
                    )
                    n += 1
        if not n:
            raise LayoutError("no sales rows found")
        conn.execute(SALES_NORMALIZE, (sid,))
    print(f"{path.name}: {n} rows loaded")
    return n, skipped


def load_sales(conn):
    ingest_sales(conn, DATA / "sales_volumes.xlsx")


def parse_cn_eu(ws):
    """Rows: question label row, then a 'Mean' row with one value per model column."""
    countries, col_country, cur = {}, {}, None
    for c in range(2, ws.max_column + 1):
        cur = txt(ws.cell(2, c).value) or cur
        col_country[c] = cur
    models = {c: txt(ws.cell(3, c).value) for c in range(2, ws.max_column + 1) if txt(ws.cell(3, c).value)}
    question = txt(ws.cell(3, 1).value)
    for r in range(4, ws.max_row + 1):
        label = txt(ws.cell(r, 1).value)
        if label == "Mean":
            for c, model in models.items():
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)):
                    yield ws.cell(r, c).coordinate, col_country[c], model, question, "mean", None, v
        elif label:
            question = label


def parse_us(ws):
    """Blocks: question label row, then 'Sample total' and seven answer rows with shares per model."""
    models = {c: txt(ws.cell(1, c).value) for c in range(3, ws.max_column + 1) if txt(ws.cell(1, c).value)}
    question = None
    for r in range(1, ws.max_row + 1):
        label, kind = txt(ws.cell(r, 1).value), txt(ws.cell(r, 2).value)
        if label and not kind:
            question = label
        elif label and kind and label != "Sample total":
            for c, model in models.items():
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)):
                    yield ws.cell(r, c).coordinate, "US", model, question, "share", label, v


def ingest_studies(conn, path):
    """One customer-study workbook -> study_response. Returns (cells loaded, cells skipped)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    n = skipped = 0
    with conn.transaction():
        sid = source_id(conn, path.name)
        conn.execute("delete from public.study_response where source_id = %s", (sid,))
        for ws in wb.worksheets:
            parser = parse_cn_eu if ws.title.startswith("CN_EU") else parse_us
            for cell, market, model, question, metric, option, value in parser(ws):
                if not (market and model and question):
                    skipped += 1
                    continue
                conn.execute(
                    "insert into public.study_response (source_id, sheet, cell_ref, market, model, question, metric, answer_option, value) "
                    "values (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                    (sid, ws.title, cell, market, model, question, metric, option, value),
                )
                n += 1
        if not n:
            raise LayoutError("no study cells found")
        conn.execute(STUDY_NORMALIZE, (sid,))
    print(f"{path.name}: {n} cells loaded")
    return n, skipped


def load_studies(conn):
    ingest_studies(conn, next(DATA.glob("*customer_studies.xlsx")))


if __name__ == "__main__":
    what = set(sys.argv[1:]) or {"feedback", "sales", "studies"}
    # autocommit so each `with conn.transaction()` block is a real, committed transaction
    conn = psycopg.connect(os.environ["SUPABASE_DB_URL"], autocommit=True)
    if "sales" in what:
        load_sales(conn)
    if "studies" in what:
        load_studies(conn)
    if "feedback" in what:
        load_feedback(conn, OpenAI())
    conn.close()
