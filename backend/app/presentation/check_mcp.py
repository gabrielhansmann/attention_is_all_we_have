"""Verify the real PowerPoint MCP protocol and file operations without editing the pitch."""
from __future__ import annotations

import asyncio
import json
import sys
import tempfile
from pathlib import Path

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[3]
PITCH_DIR = ROOT / "docs/pitch"
DECK = PITCH_DIR / "bmw-product-intelligence-mcp.pptx"


def fingerprint(path: Path) -> tuple:
    """Compare slide XML, notes, relationships and geometry after a save/reopen."""
    presentation = Presentation(path)
    slides = tuple(
        (
            slide._element.xml,
            slide.notes_slide._element.xml,
            tuple(sorted((rel.reltype, rel.target_ref) for rel in slide.part.rels.values())),
        )
        for slide in presentation.slides
    )
    return presentation.slide_width, presentation.slide_height, slides


async def call(session: ClientSession, name: str, arguments: dict) -> dict:
    result = await session.call_tool(name, arguments)
    if result.isError:
        raise RuntimeError(f"MCP tool failed: {name}: {result.content}")
    for content in result.content:
        if content.type == "text":
            payload = json.loads(content.text)
            if isinstance(payload, dict) and payload.get("error"):
                raise RuntimeError(f"MCP tool failed: {name}: {payload['error']}")
            return payload
    raise RuntimeError(f"MCP tool returned no JSON text: {name}")


async def check() -> None:
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "ppt_mcp_server", "--transport", "stdio"],
        cwd=str(PITCH_DIR),
        env={"PPT_TEMPLATE_PATH": str(PITCH_DIR)},
    )
    async with asyncio.timeout(30):
        with tempfile.TemporaryDirectory(prefix="attention-ppt-mcp-") as temporary:
            copy = Path(temporary) / "pitch-roundtrip.pptx"
            smoke = Path(temporary) / "created-by-mcp.pptx"
            async with stdio_client(params) as (read, write):
                async with ClientSession(read, write) as session:
                    await session.initialize()
                    tool_list = await session.list_tools()
                    names = {tool.name for tool in tool_list.tools}
                    assert {"create_presentation", "open_presentation", "save_presentation", "manage_text"} <= names
                    await call(session, "open_presentation", {"file_path": str(DECK), "id": "pitch"})
                    info = await call(session, "get_presentation_info", {"presentation_id": "pitch"})
                    assert info["slide_count"] == 9, info
                    for index in range(9):
                        await call(session, "get_slide_info", {"slide_index": index, "presentation_id": "pitch"})
                    text = await call(session, "extract_presentation_text", {"presentation_id": "pitch"})
                    assert "Which customer needs" in json.dumps(text)
                    await call(session, "save_presentation", {"file_path": str(copy), "presentation_id": "pitch"})
                    assert fingerprint(DECK) == fingerprint(copy), "MCP round trip changed slide content or notes"

                    await call(session, "create_presentation", {"id": "smoke"})
                    await call(session, "add_slide", {"layout_index": 6, "presentation_id": "smoke"})
                    await call(session, "manage_text", {
                        "slide_index": 0, "operation": "add", "text": "PowerPoint MCP verified",
                        "left": 1, "top": 1, "width": 7, "height": 1,
                        "font_size": 28, "font_name": "Inter", "color": [28, 105, 212],
                        "presentation_id": "smoke",
                    })
                    await call(session, "save_presentation", {"file_path": str(smoke), "presentation_id": "smoke"})
                    await call(session, "open_presentation", {"file_path": str(smoke), "id": "reopened"})
                    created = Presentation(smoke)
                    assert len(created.slides) == 1
                    assert created.slides[0].shapes[0].text == "PowerPoint MCP verified"
                    print(f"MCP passed: {len(names)} tools discovered; all nine pitch slides inspected.")
                    print("Pitch round trip preserved slide XML, notes, relationships and geometry.")
                    print("Creating, adding formatted text, saving and reopening a new PPTX passed.")


if __name__ == "__main__":
    asyncio.run(check())
